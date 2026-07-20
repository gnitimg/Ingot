from __future__ import annotations

import csv
import io
import json
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from bs4 import BeautifulSoup
from docx import Document
from openpyxl import load_workbook
from PIL import Image, ImageOps
from pptx import Presentation
from pypdf import PdfReader

try:
    import pymupdf
except ModuleNotFoundError:  # Allows native-text ingestion to keep working before optional OCR deps are installed.
    pymupdf = None  # type: ignore[assignment]


IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff"}
SUPPORTED_EXTENSIONS = {
    ".txt", ".md", ".markdown", ".pdf", ".docx", ".pptx", ".xlsx",
    ".csv", ".json", ".html", ".htm", *IMAGE_EXTENSIONS,
}


@dataclass(slots=True)
class TextSection:
    text: str
    page_number: int | None = None
    metadata: dict[str, Any] = field(default_factory=dict)


@dataclass(slots=True)
class OCRTarget:
    content: bytes
    mime_type: str
    page_number: int
    metadata: dict[str, Any] = field(default_factory=dict)
    fallback_text: str = ""


@dataclass(slots=True)
class ParsedDocument:
    sections: list[TextSection] = field(default_factory=list)
    ocr_targets: list[OCRTarget] = field(default_factory=list)
    page_count: int = 0
    warnings: list[str] = field(default_factory=list)


def _native_metadata(**extra: Any) -> dict[str, Any]:
    return {"extraction_method": "native", **extra}


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "gbk"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _parse_pdf(path: Path, min_text_chars: int, max_ocr_pages: int, render_dpi: int) -> ParsedDocument:
    reader = PdfReader(path)
    result = ParsedDocument(page_count=len(reader.pages))
    render_document: Any = None
    try:
        for index, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            meaningful_chars = len(re.sub(r"\s+", "", text))
            if meaningful_chars >= min_text_chars:
                result.sections.append(
                    TextSection(text=text, page_number=index, metadata=_native_metadata(page=index))
                )
                continue
            if len(result.ocr_targets) >= max_ocr_pages:
                result.warnings.append(f"第 {index} 页疑似扫描页，但已达到 OCR_MAX_PAGES={max_ocr_pages} 限制")
                continue
            if render_document is None:
                if pymupdf is None:
                    result.warnings.append(
                        f"第 {index} 页需要 OCR，但 PyMuPDF 未安装；"
                        "请运行 pip install pymupdf"
                    )
                    continue
                render_document = pymupdf.open(path)
            pdf_page = render_document.load_page(index - 1)
            pixmap = pdf_page.get_pixmap(dpi=render_dpi, alpha=False)
            result.ocr_targets.append(
                OCRTarget(
                    content=pixmap.tobytes("png"),
                    mime_type="image/png",
                    page_number=index,
                    metadata={"page": index, "native_text_chars": meaningful_chars},
                    fallback_text=text.strip(),
                )
            )
    finally:
        if render_document is not None:
            render_document.close()
    return result


def _parse_image(path: Path) -> ParsedDocument:
    with Image.open(path) as image:
        normalized = ImageOps.exif_transpose(image).convert("RGB")
        buffer = io.BytesIO()
        normalized.save(buffer, format="PNG", optimize=True)
    return ParsedDocument(
        ocr_targets=[
            OCRTarget(
                content=buffer.getvalue(),
                mime_type="image/png",
                page_number=1,
                metadata={"source_format": path.suffix.lower().lstrip(".")},
            )
        ],
        page_count=1,
    )


def _parse_docx(path: Path) -> list[TextSection]:
    document = Document(path)
    blocks: list[str] = []
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            blocks.append(paragraph.text.strip())
    for table_index, table in enumerate(document.tables, start=1):
        rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
        blocks.append(f"表格 {table_index}\n" + "\n".join(rows))
    return [TextSection(text="\n\n".join(blocks), metadata=_native_metadata(format="docx"))]


def _parse_pptx(path: Path) -> list[TextSection]:
    presentation = Presentation(path)
    sections: list[TextSection] = []
    for index, slide in enumerate(presentation.slides, start=1):
        texts: list[str] = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                texts.append(shape.text.strip())
        if texts:
            sections.append(
                TextSection(
                    text="\n".join(texts),
                    page_number=index,
                    metadata=_native_metadata(slide=index),
                )
            )
    return sections


def _parse_xlsx(path: Path) -> list[TextSection]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sections: list[TextSection] = []
    try:
        for sheet in workbook.worksheets:
            lines: list[str] = []
            for row in sheet.iter_rows(values_only=True):
                values = [str(value) if value is not None else "" for value in row]
                if any(value.strip() for value in values):
                    lines.append(" | ".join(values))
            if lines:
                sections.append(
                    TextSection(text="\n".join(lines), metadata=_native_metadata(sheet=sheet.title))
                )
    finally:
        workbook.close()
    return sections


def _parse_csv(data: bytes) -> list[TextSection]:
    rows = csv.reader(io.StringIO(_decode_text(data)))
    lines = [" | ".join(cell.strip() for cell in row) for row in rows]
    return [TextSection(text="\n".join(lines), metadata=_native_metadata(format="csv"))]


def _parse_json(data: bytes) -> list[TextSection]:
    value = json.loads(_decode_text(data))
    return [
        TextSection(
            text=json.dumps(value, ensure_ascii=False, indent=2),
            metadata=_native_metadata(format="json"),
        )
    ]


def _parse_html(data: bytes) -> list[TextSection]:
    soup = BeautifulSoup(_decode_text(data), "html.parser")
    for element in soup(["script", "style", "noscript"]):
        element.decompose()
    title = soup.title.string.strip() if soup.title and soup.title.string else ""
    metadata = _native_metadata(format="html")
    if title:
        metadata["title"] = title
    return [TextSection(text=soup.get_text("\n", strip=True), metadata=metadata)]


def parse_document(
    path: Path,
    original_filename: str,
    *,
    ocr_min_text_chars: int = 80,
    ocr_max_pages: int = 100,
    ocr_render_dpi: int = 144,
) -> ParsedDocument:
    extension = Path(original_filename).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise ValueError(f"不支持 {extension or '无扩展名'} 文件；支持：{supported}")

    if extension == ".pdf":
        return _parse_pdf(path, ocr_min_text_chars, ocr_max_pages, ocr_render_dpi)
    if extension in IMAGE_EXTENSIONS:
        return _parse_image(path)
    if extension == ".docx":
        sections = _parse_docx(path)
    elif extension == ".pptx":
        sections = _parse_pptx(path)
    elif extension == ".xlsx":
        sections = _parse_xlsx(path)
    else:
        data = path.read_bytes()
        if extension == ".csv":
            sections = _parse_csv(data)
        elif extension == ".json":
            sections = _parse_json(data)
        elif extension in {".html", ".htm"}:
            sections = _parse_html(data)
        else:
            sections = [
                TextSection(
                    text=_decode_text(data),
                    metadata=_native_metadata(format=extension.lstrip(".")),
                )
            ]
    sections = [section for section in sections if section.text.strip()]
    return ParsedDocument(sections=sections, page_count=max(1, len(sections)))
